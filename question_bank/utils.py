"""Bulk import helper for the Question Bank Module.

Accepts an .xlsx or .csv file with columns:
question_text, question_type, options, answer, marks, difficulty
"""
import csv
import io
import json

import openpyxl

from .models import Question


REQUIRED_COLUMNS = ['question_text', 'question_type', 'answer', 'marks', 'difficulty']


def _rows_from_csv(file_obj):
    text = io.TextIOWrapper(file_obj, encoding='utf-8')
    reader = csv.DictReader(text)
    return list(reader)


def _rows_from_xlsx(file_obj):
    wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
    sheet = wb.active
    headers = [str(c.value).strip().lower() if c.value else '' for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, row))
        rows.append(record)
    return rows


def bulk_import_questions(uploaded_file, chapter, created_by):
    """
    Parses the uploaded file and bulk-creates Question rows for the given chapter.
    Returns (created_count, error_list).
    """
    filename = uploaded_file.name.lower()
    if filename.endswith('.csv'):
        rows = _rows_from_csv(uploaded_file.file)
    elif filename.endswith('.xlsx'):
        rows = _rows_from_xlsx(uploaded_file.file)
    else:
        return 0, ["Unsupported file type. Please upload a .csv or .xlsx file."]

    created = 0
    errors = []
    to_create = []

    for idx, row in enumerate(rows, start=2):  # row 1 is header
        missing = [c for c in REQUIRED_COLUMNS if not row.get(c)]
        if missing:
            errors.append(f"Row {idx}: missing required column(s) {missing}")
            continue

        options_raw = row.get('options')
        options = None
        if options_raw:
            try:
                options = json.loads(options_raw) if isinstance(options_raw, str) else options_raw
            except (json.JSONDecodeError, TypeError):
                errors.append(f"Row {idx}: 'options' is not valid JSON, skipped options.")

        try:
            marks = int(row['marks'])
        except (ValueError, TypeError):
            errors.append(f"Row {idx}: invalid 'marks' value, skipped row.")
            continue

        q_type = str(row['question_type']).strip().upper()
        difficulty = str(row['difficulty']).strip().upper()
        if q_type not in Question.QuestionType.values:
            errors.append(f"Row {idx}: invalid question_type '{q_type}', skipped row.")
            continue
        if difficulty not in Question.Difficulty.values:
            errors.append(f"Row {idx}: invalid difficulty '{difficulty}', skipped row.")
            continue

        to_create.append(Question(
            chapter=chapter,
            question_text=str(row['question_text']).strip(),
            question_type=q_type,
            options=options,
            answer=str(row['answer']).strip(),
            marks=marks,
            difficulty=difficulty,
            created_by=created_by,
        ))

    if to_create:
        Question.objects.bulk_create(to_create)
        created = len(to_create)

    return created, errors
